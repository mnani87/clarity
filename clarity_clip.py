# Clarity Clip - A Distraction-Free Clipboard for Mac.
# ClarityOne: Clear, simple, trying to do one thing well.
# Copyright (c) 2024 MCN
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License, version 3, as published by
# the Free Software Foundation. The full license is available at:
# https://www.gnu.org/licenses/gpl-3.0.html
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY REPRESENTATION OR WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE. Indeed, the program may be worse than useless. 

# clarity_clip.py

import sys
import os
import re
from datetime import datetime
import logging
import threading

from PyQt5 import QtCore, QtGui, QtWidgets
from PyQt5.QtCore import Qt, QSize, QTimer
from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QTableWidget, QTableWidgetItem,
    QPushButton, QVBoxLayout, QWidget, QMessageBox, QInputDialog,
    QLineEdit, QLabel, QHBoxLayout, QSystemTrayIcon, QMenu, QAction,
    QFileDialog
)
from PyQt5.QtGui import QIcon, QKeySequence

from bs4 import BeautifulSoup
import PyPDF2
import docx
import openpyxl

# ------------------------ Configuration ------------------------ #

def get_history_file_path():
    """Return the path to the history file in Application Support."""
    app_support = os.path.expanduser("~/Library/Application Support/ClarityClips")
    os.makedirs(app_support, exist_ok=True)
    return os.path.join(app_support, "clipboard_history_with_tags.txt")

HISTORY_FILE = get_history_file_path()

def get_log_file_path():
    """Return the path to the log file in Logs."""
    logs_dir = os.path.expanduser("~/Library/Logs/ClarityClips")
    os.makedirs(logs_dir, exist_ok=True)
    return os.path.join(logs_dir, "clipboard_manager.log")

LOG_FILE = get_log_file_path()

MAX_ENTRIES = 500
WARNING_THRESHOLD = 450

FILE_PATH_REGEX = re.compile(
    r'^(/[^/\0]*)+/\S+\.(pdf|docx|xlsx)$',
    re.IGNORECASE
)

# Initialize a lock for file operations
file_lock = threading.Lock()

# ------------------------ Logging Setup ------------------------ #

logging.basicConfig(
    filename=LOG_FILE,
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s'
)

# ------------------------ Helper Functions ------------------------ #

def extract_text_from_html(html_content):
    """Extract plain text from HTML content."""
    try:
        soup = BeautifulSoup(html_content, 'html.parser')
        return soup.get_text()
    except Exception as e:
        logging.error(f"Error extracting text from HTML: {e}")
        return None

def extract_text_from_pdf(pdf_path):
    """Extract plain text from a PDF file."""
    try:
        reader = PyPDF2.PdfReader(pdf_path)
        text = ""
        for page_num, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text()
            if page_text:
                text += page_text
            else:
                logging.warning(f"No text found on page {page_num} of {pdf_path}.")
        return text if text else None
    except Exception as e:
        logging.error(f"Error reading PDF '{pdf_path}': {e}")
        return None

def extract_text_from_word(doc_path):
    """Extract plain text from a Word document."""
    try:
        doc = docx.Document(doc_path)
        return "\n".join([para.text for para in doc.paragraphs])
    except Exception as e:
        logging.error(f"Error reading Word document '{doc_path}': {e}")
        return None

def extract_text_from_excel(excel_path):
    """Extract plain text from an Excel file."""
    try:
        workbook = openpyxl.load_workbook(excel_path, data_only=True)
        text = ""
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = " ".join([str(cell) for cell in row if cell is not None])
                if row_text:
                    text += row_text + "\n"
        return text if text else None
    except Exception as e:
        logging.error(f"Error reading Excel file '{excel_path}': {e}")
        return None

def sanitize_content(content):
    """Sanitize content by removing newline and carriage return characters."""
    # Replace ' | ' in content to prevent split issues
    return content.replace('\n', ' ').replace('\r', '').replace(' | ', ' || ')

def export_history_to_file(export_path):
    """Export the clipboard history to the specified file path."""
    try:
        with file_lock:
            with open(HISTORY_FILE, 'r', encoding='utf-8') as src, open(export_path, 'w', encoding='utf-8') as dest:
                dest.writelines(src.readlines())
        logging.info(f"Clipboard history exported to {export_path}.")
        return True, ""
    except Exception as e:
        logging.error(f"Error exporting history: {e}")
        return False, str(e)

# ------------------------ GUI Application ------------------------ #

class ClipboardManagerGUI(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Clarity Clips")
        self.setGeometry(100, 100, 900, 600)

        # Central Widget
        self.central_widget = QWidget()
        self.setCentralWidget(self.central_widget)

        # Layout
        self.layout = QVBoxLayout()
        self.central_widget.setLayout(self.layout)

        # Search Bar
        self.search_layout = QHBoxLayout()
        self.search_label = QLabel("Search:")
        self.search_bar = QLineEdit()
        self.search_bar.setPlaceholderText("Enter keyword or tag...")
        self.search_bar.textChanged.connect(self.filter_table)
        self.search_layout.addWidget(self.search_label)
        self.search_layout.addWidget(self.search_bar)
        self.layout.addLayout(self.search_layout)

        # Table for clipboard history
        self.table = QTableWidget()
        self.table.setColumnCount(3)
        self.table.setHorizontalHeaderLabels(['Timestamp', 'Content Preview', 'Tags'])
        self.table.horizontalHeader().setStretchLastSection(True)
        self.table.setSelectionBehavior(QtWidgets.QAbstractItemView.SelectRows)
        self.table.setEditTriggers(QtWidgets.QAbstractItemView.NoEditTriggers)
        self.layout.addWidget(self.table)

        # Connect double-click to add tag
        self.table.doubleClicked.connect(self.add_tag)

        # Buttons
        self.button_layout = QHBoxLayout()

        self.copy_button = QPushButton("Copy Selected")
        self.copy_button.clicked.connect(self.copy_selected)
        self.button_layout.addWidget(self.copy_button)

        self.add_tag_button = QPushButton("Add Tag")
        self.add_tag_button.clicked.connect(self.add_tag)
        self.button_layout.addWidget(self.add_tag_button)

        self.delete_button = QPushButton("Delete Selected")
        self.delete_button.clicked.connect(self.delete_selected)
        self.button_layout.addWidget(self.delete_button)

        self.export_button = QPushButton("Export History")
        self.export_button.clicked.connect(self.export_history)
        self.button_layout.addWidget(self.export_button)

        self.refresh_button = QPushButton("Refresh")
        self.refresh_button.clicked.connect(self.load_history)
        self.button_layout.addWidget(self.refresh_button)

        self.clear_history_button = QPushButton("Clear History")
        self.clear_history_button.clicked.connect(self.clear_history)
        self.button_layout.addWidget(self.clear_history_button)

        self.layout.addLayout(self.button_layout)

        # Keyboard Shortcuts
        self.copy_shortcut = QtWidgets.QShortcut(QKeySequence("Ctrl+C"), self)
        self.copy_shortcut.activated.connect(self.copy_selected)

        self.add_tag_shortcut = QtWidgets.QShortcut(QKeySequence("Ctrl+T"), self)
        self.add_tag_shortcut.activated.connect(self.add_tag)

        self.delete_shortcut = QtWidgets.QShortcut(QKeySequence("Ctrl+D"), self)
        self.delete_shortcut.activated.connect(self.delete_selected)

        self.export_shortcut = QtWidgets.QShortcut(QKeySequence("Ctrl+E"), self)
        self.export_shortcut.activated.connect(self.export_history)

        self.refresh_shortcut = QtWidgets.QShortcut(QKeySequence("Ctrl+R"), self)
        self.refresh_shortcut.activated.connect(self.load_history)

        # System Tray Integration
        self.tray_icon = QSystemTrayIcon(self)
        tray_icon_path = self.get_tray_icon_path()
        if os.path.exists(tray_icon_path):
            self.tray_icon.setIcon(QIcon(tray_icon_path))
        else:
            self.tray_icon.setIcon(self.style().standardIcon(QtWidgets.QStyle.SP_ComputerIcon))

        # Set Tray Icon Tooltip to "Clarity Clips"
        self.tray_icon.setToolTip("Clarity Clips")

        tray_menu = QMenu()

        restore_action = QAction("Restore", self)
        restore_action.triggered.connect(self.show_normal)
        tray_menu.addAction(restore_action)

        clear_history_action = QAction("Clear History", self)
        clear_history_action.triggered.connect(self.clear_history)
        tray_menu.addAction(clear_history_action)

        exit_action = QAction("Exit", self)
        exit_action.triggered.connect(self.exit_app)
        tray_menu.addAction(exit_action)

        self.tray_icon.setContextMenu(tray_menu)
        self.tray_icon.activated.connect(self.on_tray_icon_activated)
        self.tray_icon.show()

        # Set Application Name to "Clarity Clips"
        QApplication.setApplicationName("Clarity Clips")

        # Load initial history
        self.load_history()

        # Initialize QClipboard
        self.clipboard = QApplication.clipboard()
        self.clipboard.dataChanged.connect(self.on_clipboard_change)

        # Flag to ignore clipboard changes initiated by the app
        self.ignore_clipboard_change = False

        # Flag to ensure warning is shown only once
        self.warning_shown = False

        # Initialize QTimer for clipboard polling
        self.clipboard_timer = QTimer()
        self.clipboard_timer.setInterval(500)  # Check every 500 ms
        self.clipboard_timer.timeout.connect(self.poll_clipboard)
        self.clipboard_timer.start()

        # Store the last clipboard content
        self.last_clipboard_content = self.get_current_clipboard_content()

    def get_tray_icon_path(self):
        """Return the path to the tray icon image."""
        # Replace 'tray_icon.png' with the path to your tray icon image
        return os.path.join(os.path.dirname(sys.argv[0]), 'tray_icon.png')

    def get_current_clipboard_content(self):
        """Retrieve and sanitize the current clipboard content."""
        mime = self.clipboard.mimeData()
        if mime.hasText():
            return sanitize_content(mime.text().strip())
        elif mime.hasUrls():
            urls = mime.urls()
            if urls:
                file_path = urls[0].toLocalFile()
                if FILE_PATH_REGEX.match(file_path):
                    processed_content = self.process_content(file_path)
                    return sanitize_content(processed_content) if processed_content else sanitize_content(file_path)
        return ""

    def poll_clipboard(self):
        """Periodically check the clipboard for changes."""
        current_content = self.get_current_clipboard_content()
        if current_content != self.last_clipboard_content:
            self.last_clipboard_content = current_content
            logging.info(f"Clipboard change detected via polling: {current_content[:50]}...")
            self.handle_clipboard_change(current_content)

    def on_clipboard_change(self):
        """Handle clipboard changes triggered by QClipboard signals."""
        if self.ignore_clipboard_change:
            # Reset the flag and ignore this change
            self.ignore_clipboard_change = False
            logging.info("Ignored clipboard change triggered by the application.")
            return

        current_content = self.get_current_clipboard_content()
        if current_content:
            logging.info(f"Clipboard changed detected via signal: {current_content[:50]}...")
            self.handle_clipboard_change(current_content)

    def handle_clipboard_change(self, content):
        """Process and record the clipboard change."""
        # Check for duplication
        if self.is_duplicate(content):
            logging.info("Duplicate clipboard entry detected. Skipping.")
            return

        processed_content = self.process_content(content)
        if processed_content:
            sanitized_content = sanitize_content(processed_content)
            current_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            entry = f"{current_time} | {sanitized_content} | Tags: \n"
            try:
                with file_lock:
                    with open(HISTORY_FILE, 'a', encoding='utf-8') as f:
                        f.write(entry)
                logging.info(f"New clipboard entry added at {current_time}.")
                self.load_history()
                self.check_entry_limit()
            except Exception as e:
                logging.error(f"Error writing to history file: {e}")

    def is_duplicate(self, content):
        """Check if the clipboard content is a duplicate."""
        try:
            with file_lock:
                if not os.path.exists(HISTORY_FILE):
                    return False
                with open(HISTORY_FILE, 'r', encoding='utf-8') as f:
                    for line in f:
                        parts = line.strip().split(' | ', 2)
                        if len(parts) < 3:
                            continue
                        _, existing_content, _ = parts
                        if existing_content == content:
                            return True
            return False
        except Exception as e:
            logging.error(f"Error checking for duplicates: {e}")
            return False

    def process_content(self, content):
        """Process clipboard content to extract plain text."""
        # Detect if content is a Unix-like file path
        if FILE_PATH_REGEX.match(content):
            file_path = content
            file_extension = os.path.splitext(file_path)[1].lower()
            if file_extension == '.pdf':
                return extract_text_from_pdf(file_path) or content
            elif file_extension == '.docx':
                return extract_text_from_word(file_path) or content
            elif file_extension == '.xlsx':
                return extract_text_from_excel(file_path) or content
            else:
                return content
        elif "<html>" in content.lower():
            extracted = extract_text_from_html(content)
            return extracted if extracted else content
        else:
            return content

    def load_history(self):
        """Load clipboard history from file into the table."""
        self.table.setRowCount(0)
        if not os.path.exists(HISTORY_FILE):
            open(HISTORY_FILE, 'w', encoding='utf-8').close()
            return

        try:
            with file_lock:
                with open(HISTORY_FILE, 'r', encoding='utf-8') as f:
                    lines = f.readlines()
        except Exception as e:
            logging.error(f"Error reading history file: {e}")
            QMessageBox.critical(self, "Error", "Failed to read history file.")
            return

        for line in reversed(lines):  # Reverse to show latest first
            parts = line.strip().split(' | ', 2)  # Limit splits to 2
            if len(parts) < 3:
                logging.warning(f"Malformed line skipped: {line.strip()}")
                continue
            timestamp, content, tags = parts
            # Validate timestamp format
            try:
                datetime.strptime(timestamp, "%Y-%m-%d %H:%M:%S")
            except ValueError:
                logging.warning(f"Invalid timestamp format skipped: {line.strip()}")
                continue
            content_preview = (content[:100] + '...') if len(content) > 100 else content
            tags = tags.replace('Tags: ', '')
            row_position = self.table.rowCount()
            self.table.insertRow(row_position)
            self.table.setItem(row_position, 0, QTableWidgetItem(timestamp))
            self.table.setItem(row_position, 1, QTableWidgetItem(content_preview))
            self.table.setItem(row_position, 2, QTableWidgetItem(tags))

        self.table.resizeColumnsToContents()
        self.table.horizontalHeader().setStretchLastSection(True)

    def filter_table(self):
        """Filter the table based on the search query."""
        query = self.search_bar.text().lower()
        for row in range(self.table.rowCount()):
            match = False
            for column in range(self.table.columnCount()):
                item = self.table.item(row, column)
                if item and query in item.text().lower():
                    match = True
                    break
            self.table.setRowHidden(row, not match)

    def copy_selected(self):
        """Copy the full content of the selected clipboard entry back to the clipboard."""
        selected_rows = self.table.selectionModel().selectedRows()
        if not selected_rows:
            QMessageBox.warning(self, "No Selection", "Please select an entry to copy.")
            return

        # Assuming single selection
        row = selected_rows[0].row()
        timestamp_item = self.table.item(row, 0)
        content_preview_item = self.table.item(row, 1)

        if not timestamp_item or not content_preview_item:
            QMessageBox.warning(self, "Invalid Selection", "Selected entry is invalid.")
            return

        timestamp = timestamp_item.text()
        content_preview = content_preview_item.text()

        # Find the full content from the history file
        try:
            with file_lock:
                with open(HISTORY_FILE, 'r', encoding='utf-8') as f:
                    lines = f.readlines()
        except Exception as e:
            logging.error(f"Error reading history file: {e}")
            QMessageBox.critical(self, "Error", "Failed to read history file.")
            return

        full_content = None
        for line in lines:
            parts = line.strip().split(' | ', 2)  # Limit splits to 2
            if len(parts) < 3:
                continue
            line_timestamp, content, tags = parts
            preview = (content[:100] + '...') if len(content) > 100 else content
            if line_timestamp == timestamp and preview == content_preview:
                full_content = content
                break

        if full_content:
            try:
                # Set flag to ignore the next clipboard change
                self.ignore_clipboard_change = True
                self.clipboard.setText(full_content)
                QMessageBox.information(self, "Copied", "Selected entry has been copied to clipboard.")
                logging.info(f"Copied entry from {timestamp}.")
            except Exception as e:
                logging.error(f"Error copying to clipboard: {e}")
                QMessageBox.critical(self, "Error", "Failed to copy to clipboard.")
        else:
            QMessageBox.warning(self, "Not Found", "Failed to find the selected entry.")

    def add_tag(self):
        """Add a tag to the selected clipboard entry."""
        selected_rows = self.table.selectionModel().selectedRows()
        if not selected_rows:
            QMessageBox.warning(self, "No Selection", "Please select an entry to add a tag.")
            return

        # Assuming single selection
        row = selected_rows[0].row()
        timestamp_item = self.table.item(row, 0)
        content_preview_item = self.table.item(row, 1)
        tags_item = self.table.item(row, 2)

        if not timestamp_item or not content_preview_item:
            QMessageBox.warning(self, "Invalid Selection", "Selected entry is invalid.")
            return

        timestamp = timestamp_item.text()
        content_preview = content_preview_item.text()
        existing_tags = tags_item.text()

        tag, ok = QInputDialog.getText(self, "Add Tag", "Enter tag(s) separated by commas:")
        if ok and tag.strip():
            new_tags = ','.join([t.strip() for t in tag.split(',')])
            updated_tags = ','.join(filter(None, [existing_tags, new_tags]))
            updated_tags = updated_tags.strip(',')

            # Update the history file
            try:
                with file_lock:
                    with open(HISTORY_FILE, 'r', encoding='utf-8') as f:
                        lines = f.readlines()
            except Exception as e:
                logging.error(f"Error reading history file: {e}")
                QMessageBox.critical(self, "Error", "Failed to read history file.")
                return

            updated = False
            with file_lock:
                with open(HISTORY_FILE, 'w', encoding='utf-8') as f:
                    for line in lines:
                        parts = line.strip().split(' | ', 2)  # Limit splits to 2
                        if len(parts) < 3:
                            f.write(line)
                            continue
                        line_timestamp, content, tags = parts
                        preview = (content[:100] + '...') if len(content) > 100 else content
                        if line_timestamp == timestamp and preview == content_preview:
                            f.write(f"{line_timestamp} | {content} | Tags: {updated_tags}\n")
                            updated = True
                        else:
                            f.write(line)

            if updated:
                logging.info(f"Added tags to entry from {timestamp}.")
                self.load_history()
                QMessageBox.information(self, "Success", "Tags added successfully.")
            else:
                QMessageBox.warning(self, "Not Found", "Failed to find the selected entry.")

    def delete_selected(self):
        """Delete the selected clipboard entry."""
        selected_rows = self.table.selectionModel().selectedRows()
        if not selected_rows:
            QMessageBox.warning(self, "No Selection", "Please select an entry to delete.")
            return

        reply = QMessageBox.question(
            self, 'Delete Entry',
            "Are you sure you want to delete the selected entry?",
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.No
        )

        if reply == QMessageBox.Yes:
            # Assuming single selection
            row = selected_rows[0].row()
            timestamp_item = self.table.item(row, 0)
            content_preview_item = self.table.item(row, 1)

            if not timestamp_item or not content_preview_item:
                QMessageBox.warning(self, "Invalid Selection", "Selected entry is invalid.")
                return

            timestamp = timestamp_item.text()
            content_preview = content_preview_item.text()

            # Remove the entry from the history file
            try:
                with file_lock:
                    with open(HISTORY_FILE, 'r', encoding='utf-8') as f:
                        lines = f.readlines()

                with file_lock:
                    with open(HISTORY_FILE, 'w', encoding='utf-8') as f:
                        for line in lines:
                            parts = line.strip().split(' | ', 2)  # Limit splits to 2
                            if len(parts) < 3:
                                f.write(line)
                                continue
                            line_timestamp, content, tags = parts
                            preview = (content[:100] + '...') if len(content) > 100 else content
                            if line_timestamp == timestamp and preview == content_preview:
                                # Skip writing this line to delete it
                                continue
                            else:
                                f.write(line)

                logging.info(f"Deleted entry from {timestamp}.")
                self.load_history()
                QMessageBox.information(self, "Success", "Selected entry has been deleted.")
            except Exception as e:
                logging.error(f"Error deleting entry: {e}")
                QMessageBox.critical(self, "Error", "Failed to delete the selected entry.")

    def export_history(self):
        """Export the clipboard history to a text file."""
        default_filename = f"clipboard_history_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt"
        options = QFileDialog.Options()
        file_path, _ = QFileDialog.getSaveFileName(
            self, "Export Clipboard History", default_filename, "Text Files (*.txt);;All Files (*)", options=options
        )
        if file_path:
            success, message = export_history_to_file(file_path)
            if success:
                QMessageBox.information(self, "Export Successful", f"Clipboard history exported to {file_path}.")
            else:
                QMessageBox.critical(self, "Export Failed", f"Failed to export clipboard history.\nError: {message}")

    def show_warning(self):
        """Show a warning message when the history is approaching the limit."""
        if not self.warning_shown:
            reply = QMessageBox.warning(
                self,
                "Clipboard History Warning",
                f"Clipboard history is reaching its limit of {MAX_ENTRIES} entries.\n"
                "Please export your clipboard history to retain older entries, "
                "as they may soon start getting deleted.",
                QMessageBox.Yes | QMessageBox.No,
                QMessageBox.Yes
            )
            if reply == QMessageBox.Yes:
                self.export_history()
            self.warning_shown = True

    def check_entry_limit(self):
        """Check if the number of entries has reached the warning threshold."""
        try:
            with file_lock:
                with open(HISTORY_FILE, 'r', encoding='utf-8') as f:
                    lines = f.readlines()
            entry_count = len(lines)
            if entry_count == WARNING_THRESHOLD:
                # Trigger warning in the GUI
                self.show_warning()
            elif entry_count > MAX_ENTRIES:
                # Remove the oldest entries to maintain the limit
                lines = lines[-MAX_ENTRIES:]
                with file_lock:
                    with open(HISTORY_FILE, 'w', encoding='utf-8') as f:
                        f.writelines(lines)
                logging.info(f"Clipboard entries trimmed to the last {MAX_ENTRIES} entries.")
        except Exception as e:
            logging.error(f"Error checking entry limit: {e}")

    def show_normal(self):
        """Restore the window from the system tray."""
        self.show()
        self.setWindowState(Qt.WindowActive)
        self.warning_shown = False  # Reset warning flag when restoring

    def exit_app(self):
        """Exit the application gracefully."""
        confirm = QMessageBox.question(
            self, 'Exit Application',
            "Are you sure you want to exit Clarity Clips?",
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.No
        )
        if confirm == QMessageBox.Yes:
            self.tray_icon.hide()
            QApplication.instance().quit()

    def on_tray_icon_activated(self, reason):
        """Handle tray icon activation."""
        if reason == QSystemTrayIcon.Trigger:
            self.show_normal()

    def closeEvent(self, event):
        """Handle the window close event to minimize to tray."""
        event.ignore()
        self.hide()
        self.tray_icon.showMessage(
            "Clarity Clips",
            "Application minimized to tray. Right-click the tray icon for options.",
            QSystemTrayIcon.Information,
            2000
        )

    def clear_history(self):
        """Clear the entire clipboard history."""
        reply = QMessageBox.question(
            self, 'Clear History',
            "Are you sure you want to clear all clipboard history?",
            QMessageBox.Yes | QMessageBox.No, QMessageBox.No
        )
        if reply == QMessageBox.Yes:
            try:
                with file_lock:
                    with open(HISTORY_FILE, 'w', encoding='utf-8') as f:
                        f.truncate(0)
                self.load_history()
                QMessageBox.information(self, "Success", "Clipboard history cleared successfully.")
                logging.info("Clipboard history cleared by user.")
            except Exception as e:
                logging.error(f"Error clearing history file: {e}")
                QMessageBox.critical(self, "Error", "Failed to clear clipboard history.")

# ------------------------ Main Function ------------------------ #

def main():
    app = QApplication(sys.argv)
    app.setApplicationName("Clarity Clips")
    app.setQuitOnLastWindowClosed(False)  # Ensure the app keeps running after window is closed
    window = ClipboardManagerGUI()
    window.show()
    sys.exit(app.exec_())

if __name__ == "__main__":
    main()
